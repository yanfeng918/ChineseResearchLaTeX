"""真实 LibreOffice/官方 XeLaTeX 集成；使用合成夹具，不调用生图服务。"""
from pathlib import Path
import json
import shutil
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from test_pipeline import PipelineTests, a, p, TEST_ROOT, REPO


def main():
    fixture = PipelineTests()
    fixture.setUp()
    # 保存本次集成结果供人工查看；不写真实 projects/。
    output = TEST_ROOT / '_artifacts'
    output.mkdir(exist_ok=True)
    root = fixture.root
    try:
        # 使用实际青年模板配置，验证中文与公共包构建入口。
        shutil.copyfile(REPO / 'projects/NSFC_Young/extraTex/@config.tex', root / 'extraTex/@config.tex')
        fixture.write('main.tex', '\\documentclass[12pt,UTF8,a4paper,fontset=none]{ctexart}\n'
                      '\\input{extraTex/@config.tex}\n\\begin{document}\n'
                      '\\input{extraTex/body.tex}\n\\input{extraTex/foundation.tex}\n'
                      '\\input{extraTex/statements.tex}\n\\end{document}\n')
        fixture.complete_text()
        selected = fixture.plan()
        # 接管文字构建证据：真正调用官方渲染器，而非使用单元测试的合成 PDF。
        first = p.build(root, fixture.state, 'text')
        text_log = (root / '.latex-cache/main.log').read_text(errors='replace')
        (output / 'text-main.log').write_text(text_log)
        fixture.state['stages']['14_compile']['status'] = 'in_progress'
        p.finish(root, fixture.state, '14_compile')
        for stage in ('15_figure_plan', '16_figure_prompts'):
            fixture.state['stages'][stage]['status'] = 'in_progress'
            p.finish(root, fixture.state, stage)
        for item in selected:
            fixture.record_image(item)
        fixture.state['stages']['17_imagegen']['status'] = 'in_progress'
        p.finish(root, fixture.state, '17_imagegen')

        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(9), Inches(4)
        for item in selected:
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1), Inches(3), Inches(1.2))
            box.text = '测试输入 ' + item['id']
            box.text_frame.paragraphs[0].font.size = Pt(22)
            box2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5), Inches(1), Inches(3), Inches(1.2))
            box2.text = '测试输出（合成夹具）'
            box2.text_frame.paragraphs[0].font.size = Pt(22)
            slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(3.7), Inches(1.3), Inches(1), Inches(0.6))
        deck_path = root / p.CONFIG['figures']['deck']
        deck.save(deck_path)
        a.check_deck(deck_path, len(selected))
        fixture.write('review/deck-fixture.json', json.dumps({'passed': True, 'checks': dict.fromkeys(p.DECK_CHECKS, True),
                                                            'notes': 'Synthetic native-object deck; not an editppt or imagegen acceptance claim.'}))
        proof = a.evidence(root, 'review/deck-fixture.json', p.DECK_CHECKS)
        fixture.state['deck'] = {'output': a.sha(deck_path), 'evidence': proof,
                                'images': {i['id']: a.sha(root / i['png']) for i in selected}}
        fixture.state['stages']['18_editable_ppt']['status'] = 'in_progress'
        p.finish(root, fixture.state, '18_editable_ppt')
        export = p.export_pdf(root, fixture.state)
        fixture.state['stages']['19_figure_pdf']['status'] = 'in_progress'
        p.finish(root, fixture.state, '19_figure_pdf')
        a.insertion(root, selected, True)
        fixture.write('docs/figure_insertion.md', 'Five synthetic figures mapped to body.tex for integration testing.')
        fixture.state['stages']['20_insert_figures']['status'] = 'in_progress'
        p.finish(root, fixture.state, '20_insert_figures')
        final = p.build(root, fixture.state, 'final')
        final_log = (root / '.latex-cache/main.log').read_text(errors='replace')
        warnings = lambda log: sorted({line.strip() for line in log.splitlines()
                                      if 'Warning:' in line or line.startswith(('Overfull', 'Underfull'))})
        import fitz
        with fitz.open(root / 'main.pdf') as doc:
            assert not any('??' in page.get_text() for page in doc), 'unresolved figure references'
            pages = {i['id']: [n + 1 for n, page in enumerate(doc) if ('Synthetic figure ' + i['id']) in page.get_text()]
                     for i in selected}
            assert all(pages.values()), 'not every caption is present in final PDF'
            # 记录实际页码；视觉验收由运行后查看页面完成，先保留为待查看。
            for number in sorted({n for values in pages.values() for n in values}):
                doc[number - 1].get_pixmap(matrix=fitz.Matrix(1, 1)).save(output / f'page-{number}.png')
        fixture.write('review/图文联合检查报告.md', 'Synthetic integration fixture: all five figure captions and references compiled; see recorded page previews.')
        fixture.write('review/图文联合检查证据.json', json.dumps({'passed': True,
                      'checks': dict.fromkeys(p.JOINT_CHECKS, True), 'input_sha256': final['input'],
                      'pdf_sha256': final['output'], 'figure_pages': pages,
                      'notes': 'Synthetic test fixture for deterministic gate behavior; not a proposal QC result.'}))
        fixture.state['stages']['21_joint_review']['status'] = 'in_progress'
        p.finish(root, fixture.state, '21_joint_review')
        fixture.write('review/最终图文编译检查报告.md', 'Synthetic fixture compiled through the official renderer with zero errors.')
        fixture.state['stages']['22_final_compile']['status'] = 'in_progress'
        p.finish(root, fixture.state, '22_final_compile')
        result = p.readiness(root, fixture.state)
        assert result['illustrated_body_ready'], result
        assert not result['submission_ready'], result
        for name in ('main.pdf', p.CONFIG['figures']['deck']):
            shutil.copyfile(root / name, output / Path(name).name)
        shutil.copyfile(root / '.latex-cache/main.log', output / 'main.log')
        # 修改图件后必须使最终图文状态失效。
        (root / selected[0]['pdf']).unlink()
        fixture.pdf(selected[0]['pdf'], 'Replacement fixture')
        invalid = p.readiness(root, fixture.state)
        assert not invalid['illustrated_body_ready']
        report = {'text_build_pages': first['pages'], 'final_build_pages': final['pages'],
                  'exported': export, 'figure_pages': pages, 'readiness_before_replacement': result,
                  'replacement_invalidated': not invalid['illustrated_body_ready'],
                  'existing_warnings': warnings(text_log),
                  'new_warnings': sorted(set(warnings(final_log)) - set(warnings(text_log))),
                  'live_imagegen_tested': False, 'editppt_reconstruction_tested': False}
        (output / 'integration.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
        print(json.dumps(report, ensure_ascii=False, indent=2))
    finally:
        fixture.doCleanups()


if __name__ == '__main__':
    main()
